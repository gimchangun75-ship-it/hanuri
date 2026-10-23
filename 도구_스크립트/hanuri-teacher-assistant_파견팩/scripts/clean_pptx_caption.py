# -*- coding: utf-8 -*-
"""
[hanuri-teacher-assistant] clean_pptx_caption.py (V2.0.0 Production Master)
CRIT-1/2, H-3, M-1, M-4, L-2 전량 반영 최종 완제판.
- CRIT-1 : image_part._blob 비공개 API 완전 제거 (ZIP 재팩 대신 슬라이드 전용 이미지 파트 교체)
- CRIT-2 : patch_box 지정 시 텍스트 상자 유무 무관 이미지 패치 엔진 최우선
- H-3    : int 픽셀 / float 비율 명확 판별 + x1<x2, y1<y2 clamp 보증
- M-1    : modified=False 이면 prs.save 생략 조기 반환
- M-4    : blip.get(qn('r:embed')) 사용 (lxml 공개 API)
- L-2    : os.replace 실패 시 shutil.move 폴백
- FIX-HIGH-5 : 공유 이미지 파트를 직접 고치지 않고 이 도형 전용 이미지로 교체 (ZIP 재팩 제거)
- FIX-HIGH-6 : 통짜 이미지는 box 좌표 없이 고정 영역을 칠하지 않음
"""

import sys
import os
import shutil
import re
import tempfile
import copy
import io
import xml.sax.saxutils


def safe_print(*args, **kwargs):
    if sys.stdout is not None:
        try:
            print(*args, **kwargs)
        except Exception:
            pass


if sys.stdout is not None:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.oxml.ns import qn, nsdecls
    from pptx.oxml import parse_xml
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    safe_print("python-pptx 모듈이 필요합니다. 자동 설치 중...")
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.oxml.ns import qn, nsdecls
        from pptx.oxml import parse_xml
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except subprocess.CalledProcessError:
        safe_print("❌ [네트워크 오류] python-pptx 자동 설치 실패.")
        sys.exit(1)

try:
    from PIL import Image, ImageDraw
except ImportError:
    safe_print("Pillow 모듈이 필요합니다. 자동 설치 중...")
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "Pillow"])
        from PIL import Image, ImageDraw
    except subprocess.CalledProcessError:
        safe_print("❌ [네트워크 오류] Pillow 자동 설치 실패.")
        sys.exit(1)


# ---------------------------------------------------------------------------
# 유틸리티
# ---------------------------------------------------------------------------

def extract_slide_number(val):
    if isinstance(val, int):
        return val
    s = str(val).strip()
    m = re.search(r'\d+', s)
    if m:
        return int(m.group())
    raise ValueError(f"슬라이드 번호를 인식할 수 없습니다: {val}")


def is_header_or_title(shape, target_slide):
    """제목 상자나 상단 헤더 배너인지 판정하여 도형 삭제 방지."""
    try:
        if target_slide.shapes.title and shape == target_slide.shapes.title:
            return True
        if hasattr(shape, 'name') and any(
            kw in shape.name.lower() for kw in ['title', '제목', 'header', '헤더']
        ):
            return True
    except Exception:
        pass
    try:
        if shape.top < Inches(1.4) and shape.width > Inches(4.5):
            return True
    except Exception:
        pass
    try:
        if shape.top < Inches(2.0) and shape.width > Inches(3.0) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if hasattr(p, 'font') and p.font and p.font.size and p.font.size.pt >= 18:
                    return True
                for run in p.runs:
                    if hasattr(run, 'font') and run.font and run.font.size and run.font.size.pt >= 18:
                        return True
    except Exception:
        pass
    return False


def normalize_spaces(text):
    if not text:
        return ""
    t = text.replace('\u00a0', ' ').replace('\u3000', ' ')
    return re.sub(r'  +', ' ', t)


# ---------------------------------------------------------------------------
# [DUAL-1] 텍스트 핀포인트 삭제 엔진
# ---------------------------------------------------------------------------

def clean_text_in_shape(shape, keyword):
    """
    서식 보존 및 빈 XML 노드 제거.
    [FIX-MED-12] 키워드가 여러 런(run)에 걸쳐 있어도 각 런에서 해당 글자만 지워
    나머지 글자의 서식(굵게·색 등)을 그대로 유지한다.
    """
    replaced_count = 0
    if not shape.has_text_frame:
        return replaced_count

    for p in shape.text_frame.paragraphs:
        while keyword in ''.join(r.text for r in p.runs):
            runs = list(p.runs)
            full = ''.join(r.text for r in runs)
            start = full.index(keyword)
            end = start + len(keyword)
            offset = 0
            for r in runs:
                r_start, r_end = offset, offset + len(r.text)
                offset = r_end
                cut_s, cut_e = max(start, r_start), min(end, r_end)
                if cut_s < cut_e:
                    r.text = r.text[:cut_s - r_start] + r.text[cut_e - r_start:]
            for r in runs:
                if not r.text:
                    try:
                        p._p.remove(r._r)
                    except Exception:
                        pass
            replaced_count += 1

    return replaced_count


def search_and_remove_in_shapes(shapes_collection, sp_tree, keywords, target_slide):
    deleted_shapes = 0
    cleaned_texts = 0
    shapes_to_remove = []

    for shape in shapes_collection:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            d_s, c_t = search_and_remove_in_shapes(
                shape.shapes, shape._element, keywords, target_slide
            )
            deleted_shapes += d_s
            cleaned_texts += c_t
            continue

        if shape.has_text_frame:
            full_text = shape.text_frame.text.strip()
            if is_header_or_title(shape, target_slide) or len(full_text) > 50:
                for kw in keywords:
                    kw = kw.strip()
                    if kw and kw in full_text:
                        cnt = clean_text_in_shape(shape, kw)
                        if cnt > 0:
                            cleaned_texts += cnt
                            safe_print(f"     -> [텍스트 정제] 제목/긴본문 보존 및 키워드 제거: '{kw}'")
                continue

            for kw in keywords:
                kw = kw.strip()
                if not kw:
                    continue
                if kw in full_text:
                    if len(full_text) <= len(kw) + 10:
                        shapes_to_remove.append(shape)
                        deleted_shapes += 1
                        safe_print(f"     -> [도형 삭제 예약] 캡션 글상자 식별: '{full_text[:30]}'")
                        break
                    else:
                        cnt = clean_text_in_shape(shape, kw)
                        if cnt > 0:
                            cleaned_texts += cnt
                            safe_print(f"     -> [텍스트 정제] 복합 글상자 키워드 제거: '{kw}'")

    for shape in shapes_to_remove:
        try:
            sp_tree.remove(shape._element)
        except Exception:
            try:
                shape.text_frame.text = ""
            except Exception:
                pass

    return deleted_shapes, cleaned_texts


# ---------------------------------------------------------------------------
# [DUAL-2] 이미지 패치 엔진 (CRIT-1/2, H-3, M-4)
# ---------------------------------------------------------------------------

def _parse_box_coords(box_coords, img_w, img_h):
    """
    H-3: int 픽셀 vs float 비율 명확 판별 + x1<x2, y1<y2 clamp 보증.
    모든 값이 float이고 [0.0, 1.0] 범위이면 비율 모드, 그 외 픽셀 모드.
    """
    bx, by, bw, bh = box_coords
    if all(isinstance(v, float) and 0.0 <= v <= 1.0 for v in (bx, by, bw, bh)):
        x1 = int(bx * img_w)
        y1 = int(by * img_h)
        x2 = int((bx + bw) * img_w)
        y2 = int((by + bh) * img_h)
    else:
        x1, y1 = int(bx), int(by)
        x2, y2 = int(bx + bw), int(by + bh)

    # clamp: 이미지 경계 + x1<x2, y1<y2 보증
    x1 = max(0, min(x1, img_w - 1))
    y1 = max(0, min(y1, img_h - 1))
    x2 = max(x1 + 1, min(x2, img_w))
    y2 = max(y1 + 1, min(y2, img_h))
    return x1, y1, x2, y2


def patch_image_shape(slide, box_coords=None, patch_img_path=None, fill_color=None):
    """
    [DUAL-2 / CRIT-1] 통짜 이미지(PICTURE) 슬라이드 인페인팅/패치 엔진.
    비공개 API(_blob 직접 쓰기) 완전 제거.
    반환값: 성공 여부 (bool)
    """
    picture_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not picture_shapes:
        safe_print("   ⚠️ 슬라이드에 PICTURE 도형이 없습니다.")
        return False

    pic = picture_shapes[0]
    try:
        # M-4: lxml 공개 API blip.get(qn('r:embed')) 사용
        blip_fill = pic._element.find(qn('p:blipFill'))
        if blip_fill is None:
            safe_print("   ❌ p:blipFill 요소를 찾을 수 없습니다.")
            return False
        blip = blip_fill.find(qn('a:blip'))
        if blip is None:
            safe_print("   ❌ a:blip 요소를 찾을 수 없습니다.")
            return False

        rId = blip.get(qn('r:embed'))  # M-4 fix
        if not rId:
            safe_print("   ❌ r:embed 속성이 없습니다.")
            return False

        image_part = pic.part.related_part(rId)

        orig_img = Image.open(io.BytesIO(image_part.blob)).convert('RGBA')
        img_w, img_h = orig_img.size

        # [FIX-HIGH-6] 좌표 없이 고정 영역을 칠하면 지도 등 원본이 훼손되므로 중단
        if not box_coords:
            safe_print("   ❌ 통짜 이미지는 box:x,y,w,h 좌표가 필요합니다 (예: box:0.05,0.08,0.85,0.09)")
            return False

        # H-3: 좌표 변환 + clamp
        x1, y1, x2, y2 = _parse_box_coords(box_coords, img_w, img_h)

        # 패치 적용
        if patch_img_path and os.path.exists(patch_img_path):
            safe_print(f"   -> [이미지 패치] 외부 패치 이미지 덮어쓰기: {patch_img_path}")
            patch_img = Image.open(patch_img_path).convert('RGBA')
            patch_resized = patch_img.resize((x2 - x1, y2 - y1), Image.Resampling.LANCZOS)
            orig_img.paste(patch_resized, (x1, y1), patch_resized)
        else:
            # 주변 픽셀 평균으로 인페인팅
            if fill_color is None:
                sample_pixels = []
                step_x = max(1, (x2 - x1) // 20)
                step_y = max(1, (y2 - y1) // 20)
                for px in range(x1, x2, step_x):
                    if y1 > 0:
                        sample_pixels.append(orig_img.getpixel((px, max(0, y1 - 2))))
                    if y2 < img_h:
                        sample_pixels.append(orig_img.getpixel((px, min(img_h - 1, y2 + 1))))
                for py in range(y1, y2, step_y):
                    if x1 > 0:
                        sample_pixels.append(orig_img.getpixel((max(0, x1 - 2), py)))
                    if x2 < img_w:
                        sample_pixels.append(orig_img.getpixel((min(img_w - 1, x2 + 1), py)))
                if sample_pixels:
                    avg_r = sum(p[0] for p in sample_pixels) // len(sample_pixels)
                    avg_g = sum(p[1] for p in sample_pixels) // len(sample_pixels)
                    avg_b = sum(p[2] for p in sample_pixels) // len(sample_pixels)
                    fill_color = (avg_r, avg_g, avg_b, 255)
                else:
                    fill_color = (255, 255, 255, 255)

            safe_print(
                f"   -> [이미지 인페인팅] ({x1},{y1})~({x2},{y2}) "
                f"배경색 {fill_color[:3]} 채움"
            )
            draw = ImageDraw.Draw(orig_img)
            draw.rectangle([x1, y1, x2, y2], fill=fill_color)

        # [FIX-HIGH-5] 원본 이미지 파트는 다른 슬라이드와 공유될 수 있으므로 건드리지 않는다.
        # 이 도형만 새 이미지 파트를 가리키게 하고, 더 이상 쓰이지 않는 기존 관계를 끊는다.
        out_buf = io.BytesIO()
        orig_img.save(out_buf, format='PNG')
        out_buf.seek(0)
        _, new_rId = slide.part.get_or_add_image_part(out_buf)
        blip.set(qn('r:embed'), new_rId)
        slide.part.drop_rel(rId)
        safe_print("   ✅ 이미지 패치 완료 (이 슬라이드 전용 이미지로 교체)")
        return True

    except Exception as e:
        safe_print(f"   ❌ 이미지 패치 중 오류: {e}")
        return False


# ---------------------------------------------------------------------------
# 원자적 저장 (L-2 shutil.move 폴백)
# ---------------------------------------------------------------------------

def _atomic_save(prs, pptx_path):
    """prs.save → tempfile → os.replace (L-2: shutil.move 폴백)."""
    target_dir = os.path.dirname(os.path.abspath(pptx_path))
    temp1 = None
    src = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pptx", dir=target_dir, delete=False) as tf:
            temp1 = tf.name
        prs.save(temp1)
        src = temp1
        temp1 = None

        # L-2: os.replace 실패 시 shutil.move 폴백
        try:
            os.replace(src, pptx_path)
        except OSError:
            shutil.move(src, pptx_path)
        src = None
        safe_print(f"✅ [원자적 저장 완료] 슬라이드 무손실 정제 완료: {pptx_path}")
        return True

    except PermissionError:
        safe_print(f"❌ [저장 실패 - 파일 잠김] 파워포인트를 닫고 다시 실행하세요: {pptx_path}")
        return False
    except Exception as e:
        safe_print(f"❌ 원자적 저장 실패: {e}")
        return False
    finally:
        for t in [temp1, src]:
            if t and os.path.exists(t):
                try:
                    os.remove(t)
                except Exception:
                    pass


# ---------------------------------------------------------------------------
# 진입점 함수
# ---------------------------------------------------------------------------

def clean_pptx_slide(
    pptx_path, slide_idx_1based,
    delete_keywords=None, patch_box=None, patch_image=None
):
    if not os.path.exists(pptx_path):
        safe_print(f"❌ 오류: 파일이 없습니다: {pptx_path}")
        return False

    # .bak 최초 원본 보존
    backup_path = pptx_path + ".bak"
    try:
        if not os.path.exists(backup_path):
            shutil.copy2(pptx_path, backup_path)
            safe_print(f"   (최초 원본 백업 완료: {backup_path})")
    except Exception as e:
        safe_print(f"⚠️ 백업 처리 예외: {e}")

    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        safe_print(f"❌ PPTX 로드 실패: {e}")
        return False

    total_slides = len(prs.slides)
    slide_num = extract_slide_number(slide_idx_1based)
    if slide_num < 1 or slide_num > total_slides:
        safe_print(f"❌ 슬라이드 번호 초과 (1 ~ {total_slides}): {slide_num}")
        return False

    target_slide = prs.slides[slide_num - 1]
    safe_print(f"🔍 [{slide_num}번 슬라이드 검사 시작] (전체 슬라이드: {total_slides})")

    has_text_shapes = any(
        s.has_text_frame and s.text_frame.text.strip() for s in target_slide.shapes
    )
    is_picture_slide = (
        not has_text_shapes
        and any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in target_slide.shapes)
    )

    modified = False

    # -----------------------------------------------------------------------
    # CRIT-2: patch_box 지정 시 텍스트 상자 유무와 무관하게 이미지 패치 최우선
    # -----------------------------------------------------------------------
    if patch_box is not None:
        safe_print("   ℹ️ patch_box 지정 감지 ➔ [이미지 패치 엔진 최우선 가동]")
        # 패치 실패를 '변경 없음(성공)'으로 넘기지 않도록 즉시 실패 반환
        if not patch_image_shape(target_slide, box_coords=patch_box, patch_img_path=patch_image):
            return False
        modified = True

    elif is_picture_slide:
        safe_print("   ℹ️ 통짜 이미지(PICTURE) 슬라이드 감지 ➔ [이미지 패치 엔진 가동]")
        if not patch_image_shape(target_slide, box_coords=None, patch_img_path=patch_image):
            return False
        modified = True

    else:
        safe_print("   ℹ️ 텍스트 도형 슬라이드 감지 ➔ [텍스트 핀포인트 삭제 엔진 가동]")
        if not delete_keywords:
            safe_print("   ⚠️ 삭제할 키워드가 지정되지 않았습니다.")
            return False
        if isinstance(delete_keywords, str):
            keywords = [k.strip() for k in re.split(r'[,|/\n]', delete_keywords) if k.strip()]
        else:
            keywords = [str(k).strip() for k in delete_keywords if str(k).strip()]

        d_cnt, c_cnt = search_and_remove_in_shapes(
            target_slide.shapes, target_slide.shapes._spTree, keywords, target_slide
        )
        safe_print(f"   결과: 캡션 도형 {d_cnt}개 삭제, 텍스트 정제 {c_cnt}건 완료")
        modified = (d_cnt > 0 or c_cnt > 0)

    # M-1: 변경 없으면 저장 생략 조기 반환
    if not modified:
        safe_print("   ℹ️ 변경 사항 없음. 저장을 건너뜁니다.")
        return True

    return _atomic_save(prs, pptx_path)


# ---------------------------------------------------------------------------
# CLI 진입점
# ---------------------------------------------------------------------------

def parse_cli_box(arg):
    """
    'box:x,y,w,h' 문자열을 좌표 튜플로 변환.
    [FIX-MED-9] 숫자 중 하나라도 소수점이 있으면 4개 모두 float 로 본다.
    예전에는 0,0.1,1,0.8 처럼 섞어 쓰면 픽셀로 잘못 해석되어 1픽셀 영역이 됐다.
    """
    raw = arg[len("box:"):].split(",")
    if len(raw) != 4:
        raise ValueError(f"box 좌표는 x,y,w,h 4개여야 합니다: {arg}")
    use_float = any('.' in p for p in raw)
    return tuple(float(p) if use_float else int(p) for p in raw)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        safe_print(
            "사용법: python clean_pptx_caption.py <pptx경로> <슬라이드번호> "
            "[키워드|box:x,y,w,h]\n"
            "  픽셀 좌표 예: box:40,60,644,70\n"
            "  비율 좌표 예: box:0.05,0.08,0.85,0.09"
        )
        sys.exit(1)

    p_path = sys.argv[1].strip('"\'')
    s_idx = sys.argv[2].strip()
    kw = sys.argv[3].strip() if len(sys.argv) > 3 else None

    # 박스 좌표 파싱 (box:x,y,w,h — 정수=픽셀, 소수점 포함=float 비율)
    box = None
    if kw and kw.startswith("box:"):
        try:
            box = parse_cli_box(kw)
        except ValueError as e:
            safe_print(f"❌ {e}")
            sys.exit(1)
        kw = None

    success = clean_pptx_slide(p_path, s_idx, delete_keywords=kw, patch_box=box)
    sys.exit(0 if success else 1)
