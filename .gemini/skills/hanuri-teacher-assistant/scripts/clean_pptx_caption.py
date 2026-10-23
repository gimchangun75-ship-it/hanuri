# -*- coding: utf-8 -*-
"""
[hanuri-teacher-assistant] clean_pptx_caption.py (V2.0.0 Production Master)
CRIT-1/2, H-3, M-1, M-4, L-2 전량 반영 최종 완제판.
- CRIT-1 : image_part._blob 비공개 API 완전 제거 → prs.save 후 zipfile 원자적 재팩
- CRIT-2 : patch_box 지정 시 텍스트 상자 유무 무관 이미지 패치 엔진 최우선
- H-3    : int 픽셀 / float 비율 명확 판별 + x1<x2, y1<y2 clamp 보증
- M-1    : modified=False 이면 prs.save 생략 조기 반환
- M-4    : blip.get(qn('r:embed')) 사용 (lxml 공개 API)
- L-2    : os.replace 실패 시 shutil.move 폴백
"""

import sys
import os
import shutil
import re
import tempfile
import copy
import io
import zipfile
import xml.sax.saxutils


def safe_print(*args, **kwargs):
    if sys.stdout is not None:
        try:
            print(*args, **kwargs)
        except Exception:
            pass


if sys.stdout is not None:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.oxml.ns import qn, nsdecls
    from pptx.oxml import parse_xml
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    safe_print("python-pptx 모듈이 필요합니다. 자동 설치 중...")
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.oxml.ns import qn, nsdecls
        from pptx.oxml import parse_xml
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except subprocess.CalledProcessError:
        safe_print("❌ [네트워크 오류] python-pptx 자동 설치 실패.")
        sys.exit(1)

try:
    from PIL import Image, ImageDraw
except ImportError:
    safe_print("Pillow 모듈이 필요합니다. 자동 설치 중...")
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "Pillow"])
        from PIL import Image, ImageDraw
    except subprocess.CalledProcessError:
        safe_print("❌ [네트워크 오류] Pillow 자동 설치 실패.")
        sys.exit(1)


# ---------------------------------------------------------------------------
# 유틸리티
# ---------------------------------------------------------------------------

def extract_slide_number(val):
    if isinstance(val, int):
        return val
    s = str(val).strip()
    m = re.search(r'\d+', s)
    if m:
        return int(m.group())
    raise ValueError(f"슬라이드 번호를 인식할 수 없습니다: {val}")


def is_header_or_title(shape, target_slide):
    """제목 상자나 상단 헤더 배너인지 판정하여 도형 삭제 방지."""
    try:
        if target_slide.shapes.title and shape == target_slide.shapes.title:
            return True
        if hasattr(shape, 'name') and any(
            kw in shape.name.lower() for kw in ['title', '제목', 'header', '헤더']
        ):
            return True
    except Exception:
        pass
    try:
        if shape.top < Inches(1.4) and shape.width > Inches(4.5):
            return True
    except Exception:
        pass
    try:
        if shape.top < Inches(2.0) and shape.width > Inches(3.0) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if hasattr(p, 'font') and p.font and p.font.size and p.font.size.pt >= 18:
                    return True
                for run in p.runs:
                    if hasattr(run, 'font') and run.font and run.font.size and run.font.size.pt >= 18:
                        return True
    except Exception:
        pass
    return False


def normalize_spaces(text):
    if not text:
        return ""
    t = text.replace('\u00a0', ' ').replace('\u3000', ' ')
    return re.sub(r'  +', ' ', t)


# ---------------------------------------------------------------------------
# [DUAL-1] 텍스트 핀포인트 삭제 엔진
# ---------------------------------------------------------------------------

def clean_text_in_shape(shape, keyword):
    """서식 보존 및 빈 XML 노드 제거."""
    replaced_count = 0
    if not shape.has_text_frame:
        return replaced_count

    for p in shape.text_frame.paragraphs:
        run_replaced = False
        for run in list(p.runs):
            if keyword in run.text:
                new_text = normalize_spaces(run.text.replace(keyword, ""))
                if not new_text:
                    try:
                        p._p.remove(run._r)
                    except Exception:
                        run.text = ""
                else:
                    run.text = new_text
                replaced_count += 1
                run_replaced = True

        if not run_replaced and keyword in p.text:
            new_p_text = normalize_spaces(p.text.replace(keyword, ""))
            if p.runs:
                p.runs[0].text = new_p_text
                for r in p.runs[1:]:
                    try:
                        p._p.remove(r._r)
                    except Exception:
                        r.text = ""
                replaced_count += 1
            else:
                try:
                    safe_text = xml.sax.saxutils.escape(new_p_text)
                    r_xml = f'<a:r {nsdecls("a")}><a:t>{safe_text}</a:t></a:r>'
                    r_elem = parse_xml(r_xml)
                    end_rPr = p._p.find(qn('a:endParaRPr'))
                    if end_rPr is not None:
                        rpr_copy = copy.deepcopy(end_rPr)
                        rpr_copy.tag = qn('a:rPr')
                        r_elem.insert(0, rpr_copy)
                        end_rPr.addprevious(r_elem)
                    else:
                        p._p.append(r_elem)
                    replaced_count += 1
                except Exception:
                    pass

    return replaced_count


def search_and_remove_in_shapes(shapes_collection, sp_tree, keywords, target_slide):
    deleted_shapes = 0
    cleaned_texts = 0
    shapes_to_remove = []

    for shape in shapes_collection:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            d_s, c_t = search_and_remove_in_shapes(
                shape.shapes, shape._element, keywords, target_slide
            )
            deleted_shapes += d_s
            cleaned_texts += c_t
            continue

        if shape.has_text_frame:
            full_text = shape.text_frame.text.strip()
            if is_header_or_title(shape, target_slide) or len(full_text) > 50:
                for kw in keywords:
                    kw = kw.strip()
                    if kw and kw in full_text:
                        cnt = clean_text_in_shape(shape, kw)
                        if cnt > 0:
                            cleaned_texts += cnt
                            safe_print(f"     -> [텍스트 정제] 제목/긴본문 보존 및 키워드 제거: '{kw}'")
                continue

            for kw in keywords:
                kw = kw.strip()
                if not kw:
                    continue
                if kw in full_text:
                    if len(full_text) <= len(kw) + 10:
                        shapes_to_remove.append(shape)
                        deleted_shapes += 1
                        safe_print(f"     -> [도형 삭제 예약] 캡션 글상자 식별: '{full_text[:30]}'")
                        break
                    else:
                        cnt = clean_text_in_shape(shape, kw)
                        if cnt > 0:
                            cleaned_texts += cnt
                            safe_print(f"     -> [텍스트 정제] 복합 글상자 키워드 제거: '{kw}'")

    for shape in shapes_to_remove:
        try:
            sp_tree.remove(shape._element)
        except Exception:
            try:
                shape.text_frame.text = ""
            except Exception:
                pass

    return deleted_shapes, cleaned_texts


# ---------------------------------------------------------------------------
# [DUAL-2] 이미지 패치 엔진 (CRIT-1/2, H-3, M-4)
# ---------------------------------------------------------------------------

def _parse_box_coords(box_coords, img_w, img_h):
    """
    H-3: int 픽셀 vs float 비율 명확 판별 + x1<x2, y1<y2 clamp 보증.
    모든 값이 float이고 [0.0, 1.0] 범위이면 비율 모드, 그 외 픽셀 모드.
    """
    bx, by, bw, bh = box_coords
    if all(isinstance(v, float) and 0.0 <= v <= 1.0 for v in (bx, by, bw, bh)):
        x1 = int(bx * img_w)
        y1 = int(by * img_h)
        x2 = int((bx + bw) * img_w)
        y2 = int((by + bh) * img_h)
    else:
        x1, y1 = int(bx), int(by)
        x2, y2 = int(bx + bw), int(by + bh)

    # clamp: 이미지 경계 + x1<x2, y1<y2 보증
    x1 = max(0, min(x1, img_w - 1))
    y1 = max(0, min(y1, img_h - 1))
    x2 = max(x1 + 1, min(x2, img_w))
    y2 = max(y1 + 1, min(y2, img_h))
    return x1, y1, x2, y2


def patch_image_shape(slide, box_coords=None, patch_img_path=None, fill_color=None):
    """
    [DUAL-2 / CRIT-1] 통짜 이미지(PICTURE) 슬라이드 인페인팅/패치 엔진.
    비공개 API(_blob 직접 쓰기) 완전 제거.
    반환값: (성공: bool, ZIP내부경로: str|None, 패치바이트: bytes|None)
    실제 ZIP 재팩은 호출부(_atomic_save)에서 prs.save 이후에 수행한다.
    """
    picture_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not picture_shapes:
        safe_print("   ⚠️ 슬라이드에 PICTURE 도형이 없습니다.")
        return False, None, None

    pic = picture_shapes[0]
    try:
        # M-4: lxml 공개 API blip.get(qn('r:embed')) 사용
        blip_fill = pic._element.find(qn('p:blipFill'))
        if blip_fill is None:
            safe_print("   ❌ p:blipFill 요소를 찾을 수 없습니다.")
            return False, None, None
        blip = blip_fill.find(qn('a:blip'))
        if blip is None:
            safe_print("   ❌ a:blip 요소를 찾을 수 없습니다.")
            return False, None, None

        rId = blip.get(qn('r:embed'))  # M-4 fix
        if not rId:
            safe_print("   ❌ r:embed 속성이 없습니다.")
            return False, None, None

        image_part = pic.part.related_part(rId)
        # CRIT-1: ZIP 내부 경로 확보 (공개 API partname 사용)
        zip_part_name = image_part.partname.lstrip('/')

        orig_img = Image.open(io.BytesIO(image_part.blob)).convert('RGBA')
        img_w, img_h = orig_img.size

        # H-3: 좌표 변환 + clamp
        if box_coords:
            x1, y1, x2, y2 = _parse_box_coords(box_coords, img_w, img_h)
        else:
            # 기본값: 상단 헤더 배너 표준 영역
            x1, y1 = 40, 60
            x2, y2 = min(684, img_w), min(130, img_h)

        # 패치 적용
        if patch_img_path and os.path.exists(patch_img_path):
            safe_print(f"   -> [이미지 패치] 외부 패치 이미지 덮어쓰기: {patch_img_path}")
            patch_img = Image.open(patch_img_path).convert('RGBA')
            patch_resized = patch_img.resize((x2 - x1, y2 - y1), Image.Resampling.LANCZOS)
            orig_img.paste(patch_resized, (x1, y1), patch_resized)
        else:
            # 주변 픽셀 평균으로 인페인팅
            if fill_color is None:
                sample_pixels = []
                step_x = max(1, (x2 - x1) // 20)
                step_y = max(1, (y2 - y1) // 20)
                for px in range(x1, x2, step_x):
                    if y1 > 0:
                        sample_pixels.append(orig_img.getpixel((px, max(0, y1 - 2))))
                    if y2 < img_h:
                        sample_pixels.append(orig_img.getpixel((px, min(img_h - 1, y2 + 1))))
                for py in range(y1, y2, step_y):
                    if x1 > 0:
                        sample_pixels.append(orig_img.getpixel((max(0, x1 - 2), py)))
                    if x2 < img_w:
                        sample_pixels.append(orig_img.getpixel((min(img_w - 1, x2 + 1), py)))
                if sample_pixels:
                    avg_r = sum(p[0] for p in sample_pixels) // len(sample_pixels)
                    avg_g = sum(p[1] for p in sample_pixels) // len(sample_pixels)
                    avg_b = sum(p[2] for p in sample_pixels) // len(sample_pixels)
                    fill_color = (avg_r, avg_g, avg_b, 255)
                else:
                    fill_color = (255, 255, 255, 255)

            safe_print(
                f"   -> [이미지 인페인팅] ({x1},{y1})~({x2},{y2}) "
                f"배경색 {fill_color[:3]} 채움"
            )
            draw = ImageDraw.Draw(orig_img)
            draw.rectangle([x1, y1, x2, y2], fill=fill_color)

        out_buf = io.BytesIO()
        orig_img.save(out_buf, format='PNG')
        patched_bytes = out_buf.getvalue()
        safe_print(f"   ✅ 이미지 패치 완료 (ZIP 경로: {zip_part_name})")
        return True, zip_part_name, patched_bytes

    except Exception as e:
        safe_print(f"   ❌ 이미지 패치 중 오류: {e}")
        return False, None, None


# ---------------------------------------------------------------------------
# 원자적 저장 (CRIT-1 ZIP 재팩 + L-2 shutil.move 폴백)
# ---------------------------------------------------------------------------

def _atomic_save(prs, pptx_path, zip_patch=None):
    """
    prs.save → tempfile, zip_patch 있으면 ZIP 재팩 → os.replace (L-2: shutil.move 폴백).
    zip_patch: (zip_part_name: str, patched_bytes: bytes) | None
    """
    target_dir = os.path.dirname(os.path.abspath(pptx_path))
    temp1 = None
    temp2 = None
    src = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pptx", dir=target_dir, delete=False) as tf:
            temp1 = tf.name
        prs.save(temp1)

        if zip_patch:
            # CRIT-1: zipfile 원자적 재팩으로 비공개 API 완전 대체
            zip_part_name, patched_bytes = zip_patch
            with tempfile.NamedTemporaryFile(suffix=".pptx", dir=target_dir, delete=False) as tf2:
                temp2 = tf2.name
            with zipfile.ZipFile(temp1, 'r') as zin, \
                 zipfile.ZipFile(temp2, 'w', zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename == zip_part_name:
                        zout.writestr(item, patched_bytes)
                    else:
                        zout.writestr(item, zin.read(item.filename))
            os.remove(temp1)
            temp1 = None
            src = temp2
            temp2 = None
        else:
            src = temp1
            temp1 = None

        # L-2: os.replace 실패 시 shutil.move 폴백
        try:
            os.replace(src, pptx_path)
        except OSError:
            shutil.move(src, pptx_path)
        src = None
        safe_print(f"✅ [원자적 저장 완료] 슬라이드 무손실 정제 완료: {pptx_path}")
        return True

    except PermissionError:
        safe_print(f"❌ [저장 실패 - 파일 잠김] 파워포인트를 닫고 다시 실행하세요: {pptx_path}")
        return False
    except Exception as e:
        safe_print(f"❌ 원자적 저장 실패: {e}")
        return False
    finally:
        for t in [temp1, temp2, src]:
            if t and os.path.exists(t):
                try:
                    os.remove(t)
                except Exception:
                    pass


# ---------------------------------------------------------------------------
# 진입점 함수
# ---------------------------------------------------------------------------

def clean_pptx_slide(
    pptx_path, slide_idx_1based,
    delete_keywords=None, patch_box=None, patch_image=None
):
    if not os.path.exists(pptx_path):
        safe_print(f"❌ 오류: 파일이 없습니다: {pptx_path}")
        return False

    # .bak 최초 원본 보존
    backup_path = pptx_path + ".bak"
    try:
        if not os.path.exists(backup_path):
            shutil.copy2(pptx_path, backup_path)
            safe_print(f"   (최초 원본 백업 완료: {backup_path})")
    except Exception as e:
        safe_print(f"⚠️ 백업 처리 예외: {e}")

    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        safe_print(f"❌ PPTX 로드 실패: {e}")
        return False

    total_slides = len(prs.slides)
    slide_num = extract_slide_number(slide_idx_1based)
    if slide_num < 1 or slide_num > total_slides:
        safe_print(f"❌ 슬라이드 번호 초과 (1 ~ {total_slides}): {slide_num}")
        return False

    target_slide = prs.slides[slide_num - 1]
    safe_print(f"🔍 [{slide_num}번 슬라이드 검사 시작] (전체 슬라이드: {total_slides})")

    has_text_shapes = any(
        s.has_text_frame and s.text_frame.text.strip() for s in target_slide.shapes
    )
    is_picture_slide = (
        not has_text_shapes
        and any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in target_slide.shapes)
    )

    modified = False
    zip_patch = None  # (zip_part_name, patched_bytes)

    # -----------------------------------------------------------------------
    # CRIT-2: patch_box 지정 시 텍스트 상자 유무와 무관하게 이미지 패치 최우선
    # -----------------------------------------------------------------------
    if patch_box is not None:
        safe_print("   ℹ️ patch_box 지정 감지 ➔ [이미지 패치 엔진 최우선 가동]")
        ok, zname, zbytes = patch_image_shape(
            target_slide, box_coords=patch_box, patch_img_path=patch_image
        )
        if ok:
            modified = True
            zip_patch = (zname, zbytes)

    elif is_picture_slide:
        safe_print("   ℹ️ 통짜 이미지(PICTURE) 슬라이드 감지 ➔ [이미지 패치 엔진 가동]")
        ok, zname, zbytes = patch_image_shape(
            target_slide, box_coords=None, patch_img_path=patch_image
        )
        if ok:
            modified = True
            zip_patch = (zname, zbytes)

    else:
        safe_print("   ℹ️ 텍스트 도형 슬라이드 감지 ➔ [텍스트 핀포인트 삭제 엔진 가동]")
        if not delete_keywords:
            safe_print("   ⚠️ 삭제할 키워드가 지정되지 않았습니다.")
            return False
        if isinstance(delete_keywords, str):
            keywords = [k.strip() for k in re.split(r'[,|/\n]', delete_keywords) if k.strip()]
        else:
            keywords = [str(k).strip() for k in delete_keywords if str(k).strip()]

        d_cnt, c_cnt = search_and_remove_in_shapes(
            target_slide.shapes, target_slide.shapes._spTree, keywords, target_slide
        )
        safe_print(f"   결과: 캡션 도형 {d_cnt}개 삭제, 텍스트 정제 {c_cnt}건 완료")
        modified = (d_cnt > 0 or c_cnt > 0)

    # M-1: 변경 없으면 저장 생략 조기 반환
    if not modified:
        safe_print("   ℹ️ 변경 사항 없음. 저장을 건너뜁니다.")
        return True

    return _atomic_save(prs, pptx_path, zip_patch=zip_patch)


# ---------------------------------------------------------------------------
# CLI 진입점
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    if len(sys.argv) < 3:
        safe_print(
            "사용법: python clean_pptx_caption.py <pptx경로> <슬라이드번호> "
            "[키워드|box:x,y,w,h]\n"
            "  픽셀 좌표 예: box:40,60,644,70\n"
            "  비율 좌표 예: box:0.05,0.08,0.85,0.09"
        )
        sys.exit(1)

    p_path = sys.argv[1].strip('"\'')
    s_idx = sys.argv[2].strip()
    kw = sys.argv[3].strip() if len(sys.argv) > 3 else None

    # 박스 좌표 파싱 (box:x,y,w,h — 정수=픽셀, 소수점 포함=float 비율)
    box = None
    if kw and kw.startswith("box:"):
        raw_parts = kw.replace("box:", "").split(",")
        parts = [float(p) if '.' in p else int(p) for p in raw_parts]
        if len(parts) == 4:
            box = tuple(parts)
            kw = None

    success = clean_pptx_slide(p_path, s_idx, delete_keywords=kw, patch_box=box)
    sys.exit(0 if success else 1)
