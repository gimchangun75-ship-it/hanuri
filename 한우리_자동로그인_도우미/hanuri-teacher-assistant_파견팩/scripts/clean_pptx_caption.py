# -*- coding: utf-8 -*-
"""
[hanuri-teacher-assistant] clean_pptx_caption.py (V1.8.5 Production Master)
소부장(Claude CLI) 12차 감리 전량 반영 완제본:
1. [CRIT-NEW-2] 단일 run 빈 문자열 치환 후 XML 요소(<a:r>) 자체를 p._p.remove(run._r)로 정리하여 빈 런 잔존 완전 차단
2. [WARN-12-B] 테마 상속 폰트 방어: 위치(Y<1.4)/크기(W>4.5) 및 제목 도형/플레이스홀더 최우선 평가로 헤더 보호 누락 원천 차단
3. [WARN-12-A] cross-run 서식 보존: 첫 번째 run 서식 유지 + 나머지 빈 run node XML 안전 제거
4. [INFO-12-B] MSO_SHAPE_TYPE.GROUP 공식 enum 도입으로 매직넘버(6) 탈피
5. [CRIT-11] OOXML XSD 스펙 준수: endParaRPr 직전 a:r 삽입(addprevious) 유지
6. [CRIT-9] OOXML 직속 자식 a:endParaRPr 기반 rPr 서식 복제 유지
7. [CRIT-10] xml.sax.saxutils.escape() XML 예약문자 인젝션 방어 유지
8. [BLOCK-1] tempfile + os.replace를 통한 100% 원자적(Atomic) PPTX 저장 유지
9. [H-C] .bak 최초 원본 보존 유지
10. [H-A] safe_print 헬퍼 전면 적용
"""

import sys
import os
import shutil
import re
import tempfile
import copy
import xml.sax.saxutils

def safe_print(*args, **kwargs):
    if sys.stdout is not None:
        try:
            print(*args, **kwargs)
        except Exception:
            pass

if sys.stdout is not None:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# python-pptx 자동 설치 및 네트워크 단절 친절 안내
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.oxml.ns import qn, nsdecls
    from pptx.oxml import parse_xml
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    safe_print("python-pptx 모듈이 필요합니다. 자동 설치 중...")
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.oxml.ns import qn, nsdecls
        from pptx.oxml import parse_xml
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except subprocess.CalledProcessError:
        safe_print("❌ [네트워크 오류] python-pptx 자동 설치에 실패했습니다. 인터넷 연결을 확인해 주세요.")
        sys.exit(1)


def extract_slide_number(val):
    if isinstance(val, int):
        return val
    s = str(val).strip()
    m = re.search(r'\d+', s)
    if m:
        return int(m.group())
    raise ValueError(f"슬라이드 번호를 인식할 수 없습니다: {val}")


def is_header_or_title(shape, target_slide):
    """
    [WARN-12-B 해결] 제목 상자나 상단 헤더 배너인지 판정하여 도형 삭제를 영구 방지
    위치·크기 및 슬라이드 타이틀 객체를 폰트 크기보다 최우선 평가
    """
    # 1. pptx 기본 타이틀 도형이거나 이름에 Title/제목이 포함된 경우
    try:
        if target_slide.shapes.title and shape == target_slide.shapes.title:
            return True
        if hasattr(shape, 'name') and any(kw in shape.name.lower() for kw in ['title', '제목', 'header', '헤더']):
            return True
    except Exception:
        pass

    # 2. [최우선 위치 가드] 상단 Y < 1.4인치 이고 너비 W > 4.5인치인 배너/헤더
    try:
        if shape.top < Inches(1.4) and shape.width > Inches(4.5):
            return True
    except Exception:
        pass

    # 3. [WARN-13-C & WARN-14-A 선제 반영] 폰트 크기 기반 판정:
    # 상단(Y < 2.0인치) 및 넓은 너비(W > 3.0인치)를 가진 대형 제목에 한해 보호하며, 좁은 지도 라벨의 오탐을 원천 차단
    try:
        if shape.top < Inches(2.0) and shape.width > Inches(3.0) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if hasattr(p, 'font') and p.font and p.font.size and p.font.size.pt >= 18:
                    return True
                for run in p.runs:
                    if hasattr(run, 'font') and run.font and run.font.size and run.font.size.pt >= 18:
                        return True
    except Exception:
        pass

    return False


def normalize_spaces(text):
    """
    NBSP(\\u00a0) 및 전각 공백(\\u3000)을 일반 공백으로 정규화 후 연속 공백 압축
    """
    if not text:
        return ""
    t = text.replace('\u00a0', ' ').replace('\u3000', ' ')
    return re.sub(r'  +', ' ', t)


def clean_text_in_shape(shape, keyword):
    """
    [CRIT-NEW-2 & WARN-12-A] 서식 보존 및 빈 XML 노드 제거:
    단일 run 우선 치환 후 내용이 빈 경우 XML 노드(<a:r>) 제거,
    복합 run에 걸친 경우 기존 run의 rPr 서식을 유지하면서 정제
    """
    replaced_count = 0
    if not shape.has_text_frame:
        return replaced_count

    for p in shape.text_frame.paragraphs:
        run_replaced = False
        # [CRIT-NEW-2 해결] list(p.runs) 복사본 순회로 안전한 요소 제거
        for run in list(p.runs):
            if keyword in run.text:
                new_text = normalize_spaces(run.text.replace(keyword, ""))
                if not new_text:
                    try:
                        p._p.remove(run._r)
                    except Exception:
                        run.text = ""
                else:
                    run.text = new_text
                replaced_count += 1
                run_replaced = True

        if not run_replaced and keyword in p.text:
            safe_print(f"     -> [안내] '{keyword}' 키워드가 분할 run에 걸쳐 있어 단락 통합 정제를 수행합니다.")
            new_p_text = p.text.replace(keyword, "")
            new_p_text = normalize_spaces(new_p_text)
            
            if p.runs:
                # [WARN-12-A 해결] 첫 번째 run에 텍스트를 배치하고 서식(rPr)이 있던 나머지 빈 run은 xml 트리에서 안전 정리
                p.runs[0].text = new_p_text
                runs_to_remove = p.runs[1:]
                for r in runs_to_remove:
                    try:
                        p._p.remove(r._r)
                    except Exception:
                        r.text = ""
                replaced_count += 1
            else:
                try:
                    safe_text = xml.sax.saxutils.escape(new_p_text)
                    r_xml = f'<a:r {nsdecls("a")}><a:t>{safe_text}</a:t></a:r>'
                    r_elem = parse_xml(r_xml)
                    
                    # [CRIT-11 해결] XSD 스펙 준수: endParaRPr 직전 삽입 또는 append
                    end_rPr = p._p.find(qn('a:endParaRPr'))
                    if end_rPr is not None:
                        rpr_copy = copy.deepcopy(end_rPr)
                        rpr_copy.tag = qn('a:rPr')
                        r_elem.insert(0, rpr_copy)
                        end_rPr.addprevious(r_elem)
                    else:
                        p._p.append(r_elem)
                        
                    replaced_count += 1
                except Exception as e_xml:
                    safe_print(f"     -> [경고] XML 레벨 run 생성 예외: {e_xml}")

    return replaced_count


def search_and_remove_in_shapes(shapes_collection, sp_tree, keywords, target_slide):
    deleted_shapes = 0
    cleaned_texts = 0
    shapes_to_remove = []

    for shape in shapes_collection:
        # [INFO-12-B 해결] MSO_SHAPE_TYPE.GROUP 공식 enum 사용
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            d_s, c_t = search_and_remove_in_shapes(shape.shapes, shape._element, keywords, target_slide)
            deleted_shapes += d_s
            cleaned_texts += c_t
            continue

        if shape.has_text_frame:
            full_text = shape.text_frame.text.strip()
            # 50자 초과 본문/제목 도형인 경우: 모든 키워드를 순회하여 전수 정제
            if is_header_or_title(shape, target_slide) or len(full_text) > 50:
                for kw in keywords:
                    kw = kw.strip()
                    if kw and kw in full_text:
                        cnt = clean_text_in_shape(shape, kw)
                        cleaned_texts += cnt
                        if cnt > 0:
                            safe_print(f"  -> [제목/본문 도형 100% 보존] 키워드: '{kw}' (서식 보존 텍스트 정제: {cnt}건)")
                            full_text = shape.text_frame.text.strip()
            else:
                # 50자 이하 독립 캡션 라벨 도형인 경우: 키워드 1개라도 매칭되면 도형 삭제 등록 후 break
                for kw in keywords:
                    kw = kw.strip()
                    if kw and kw in full_text:
                        shapes_to_remove.append(shape)
                        safe_print(f"  -> [지도 캡션 라벨 삭제] 키워드: '{kw}' | 내용: \"{full_text}\"")
                        break

    for shape in shapes_to_remove:
        try:
            sp_tree.remove(shape._element)
            deleted_shapes += 1
        except Exception:
            pass

    return deleted_shapes, cleaned_texts


def remove_captions_from_slide(pptx_path, slide_number_raw, keywords_to_delete):
    if not os.path.exists(pptx_path):
        safe_print(f"[오류] 파일이 존재하지 않습니다: {pptx_path}")
        return False

    try:
        slide_number = extract_slide_number(slide_number_raw)
    except ValueError as e:
        safe_print(f"[오류] {e}")
        return False

    backup_path = pptx_path + ".bak"
    try:
        if not os.path.exists(backup_path):
            shutil.copy2(pptx_path, backup_path)
            safe_print(f"[안전 최초 백업 완료] {os.path.basename(backup_path)}")
        else:
            safe_print(f"[백업 유지] 기존 최초 백업본 보존: {os.path.basename(backup_path)}")
    except Exception as e:
        safe_print(f"[경고] 백업 처리 예외: {e}")

    try:
        prs = Presentation(pptx_path)
    except PermissionError:
        safe_print(f"[오류] 파일이 PowerPoint에서 열려 있습니다. 파워포인트를 닫은 후 다시 시도해주세요.")
        return False
    except Exception as e:
        safe_print(f"[오류] 프레젠테이션 열기 실패: {e}")
        return False

    total_slides = len(prs.slides)
    if slide_number < 1 or slide_number > total_slides:
        safe_print(f"[오류] 슬라이드 번호가 범위를 벗어났습니다. (전체 {total_slides}장 중 {slide_number}번 요청됨)")
        return False

    target_slide = prs.slides[slide_number - 1]
    # NOTE: private API _spTree, python-pptx >= 0.6 기준 도형 삭제 표준 우회법
    sp_tree = target_slide.shapes._spTree

    if isinstance(keywords_to_delete, str):
        keywords = [k.strip() for k in keywords_to_delete.split(",") if k.strip()]
    else:
        keywords = [str(k).strip() for k in keywords_to_delete if str(k).strip()]

    deleted_count, cleaned_count = search_and_remove_in_shapes(target_slide.shapes, sp_tree, keywords, target_slide)

    target_dir = os.path.dirname(os.path.abspath(pptx_path))
    temp_file_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pptx", dir=target_dir, delete=False) as tf:
            temp_file_path = tf.name

        prs.save(temp_file_path)
        os.replace(temp_file_path, pptx_path)
        temp_file_path = None
        safe_print(f"✅ [원자적 저장 완료] {slide_number}번 슬라이드 무결점 교정 성공! (도형 삭제: {deleted_count}개, 서식 보존 정제: {cleaned_count}건)")
        return True
    except PermissionError:
        if temp_file_path and os.path.exists(temp_file_path):
            try:
                os.remove(temp_file_path)
            except Exception:
                pass
        safe_print(f"[오류] 저장 실패: 파일이 PowerPoint에서 열려 있습니다. 저장 후 닫아주세요.")
        return False
    except Exception as e:
        if temp_file_path and os.path.exists(temp_file_path):
            try:
                os.remove(temp_file_path)
            except Exception:
                pass
        safe_print(f"[오류] 프레젠테이션 원자적 저장 실패: {e}")
        return False


if __name__ == "__main__":
    if len(sys.argv) < 4:
        safe_print("사용법: python clean_pptx_caption.py <PPTX파일경로> <슬라이드번호> <삭제할키워드1,키워드2...>")
        safe_print("예시:   python clean_pptx_caption.py D:\\수업.pptx 9 '애버딘,북해 유전'")
        sys.exit(1)

    path = sys.argv[1].strip('"\'')
    s_num = sys.argv[2].strip('"\'')
    kws = sys.argv[3].strip('"\'')

    success = remove_captions_from_slide(path, s_num, kws)
    sys.exit(0 if success else 1)
